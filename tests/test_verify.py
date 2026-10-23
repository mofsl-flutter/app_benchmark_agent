"""Unit tests for verify.py (the RISE re-check). Run: python3 -m unittest discover tests"""
import os
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(HERE))
os.environ.setdefault("RIISE_BENCH_HOME", tempfile.mkdtemp())

import benchlib as B  # noqa: E402
import build_report as R  # noqa: E402
import verify as V  # noqa: E402


class VerdictTests(unittest.TestCase):
    def test_direction_and_noise_band(self):
        self.assertEqual(V.verdict("L1", 600, 500)[0], "better")    # lower is better
        self.assertEqual(V.verdict("S5", 50, 60)[0], "better")      # higher is better
        self.assertEqual(V.verdict("L1", 500, 600)[0], "worse")
        self.assertEqual(V.verdict("L1", 600, 590)[0], "same")      # -1.7 % is inside the band
        self.assertEqual(V.verdict("L1", None, 590)[0], "new")
        self.assertEqual(V.verdict("X4", 0, 3)[0], "worse")         # no % from zero, the sign decides

    def test_compare_scores_and_kept(self):
        before = {"RISE": {"L1": 900, "R2": 900}, "Groww": {"L1": 600, "R2": 400}, "Kite": {"L1": 700, "R2": 300}}
        res = V.compare(before, {"L1": 500}, ["RISE", "Groww", "Kite"], B.DEFAULT_WEIGHTS)
        row = res["rows"][0]
        self.assertEqual((row["before"], row["after"], row["verdict"]), (900, 500, "better"))
        self.assertEqual((row["rank_before"], row["rank_after"]), (3, 1))
        self.assertEqual(res["kept"], ["R2"])
        self.assertEqual(res["categories"]["Loading Time"], (0.0, 100.0))
        self.assertLess(res["overall"][0], res["overall"][1])
        self.assertEqual(before["RISE"]["L1"], 900)                  # the report's values are not modified


class WorkbookTests(unittest.TestCase):
    def test_parse_update_round_trip(self):
        from openpyxl import load_workbook
        import build_benchmark as W
        module = dict(B.load_module("quote"), coldstart=True)
        apps = ["RISE", "Groww", "Kite"]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "quote_dev_phone_2026-10-22.xlsx"
            layout = W.build_workbook(path, apps=apps, device_label="samsung-SM-E146B", module=module)
            wb = load_workbook(path)
            for app, vals in (("RISE", [900, 800, 1000]), ("Groww", [600, 600, 600]), ("Kite", [700, 700, 700])):
                sheet, row = layout["runs"][("L1", app)]
                for k, v in enumerate(vals):
                    wb[sheet].cell(row=row, column=R.RUN_FIRST + k).value = v
            item = module["offline_items"][0]
            wb[B.OFF].cell(row=layout["off_rows"][B.slug(item)], column=layout["off_cols"]["RISE"]).value = "No"
            wb.save(path)

            wb = load_workbook(path)
            rep = V.parse_workbook(wb)
            self.assertEqual(rep["device"], "samsung-SM-E146B")
            self.assertEqual(rep["apps"], apps)
            med = V.report_medians(rep, module)
            self.assertEqual(med["RISE"]["L1"], 900)
            self.assertEqual(med["RISE"]["O1"], 0.0)
            self.assertNotIn("L1", V.report_medians(rep, dict(module, coldstart=False))["RISE"])

            rows = [{"app": "RISE", "metric": "L1", "run": str(i), "num": float(v), "value": str(v), "timestamp": "t"}
                    for i, v in enumerate([500, 520, 510], 1)]
            rows.append({"app": "RISE", "metric": f"O1.{B.slug(item)}", "run": "1", "num": None, "value": "Yes", "timestamp": "t"})
            after = R.medians(rows, module, {"RISE"})["RISE"]
            res = V.compare(med, after, apps, B.DEFAULT_WEIGHTS)
            meta = {"date": "2026-10-23", "version_before": "6.2.09", "version_after": "6.3.00", "module": module}
            written = V.update_workbook(wb, rep, res, rows, [], meta)
            self.assertEqual(sorted(written), ["L1", "O1"])
            sheet, row = layout["runs"][("L1", "RISE")]
            self.assertEqual([wb[sheet].cell(row=row, column=R.RUN_FIRST + k).value for k in range(5)], [500, 520, 510, None, None])
            sheet, row = layout["runs"][("L1", "Groww")]
            self.assertEqual(wb[sheet].cell(row=row, column=R.RUN_FIRST).value, 600)   # competitors untouched
            self.assertIn(V.VERIFY_SHEET, wb.sheetnames)
