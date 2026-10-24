#!/usr/bin/env python3
"""Workbook layout for one module (14 tabs incl. Summary and Scorecard), as a library.

build_report.py calls build_workbook(...) and then fills the Run cells from runs.csv. Nobody types
into the generated workbook: weights come from weights.json and values from the CSVs.

Standalone: python3 build_benchmark.py out.xlsx [module]   (empty workbook, default module portfolio)
"""
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

import benchlib as B

NRUNS = B.NRUNS
F = "Arial"
f_base = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=16, bold=True, color="1F3864")
f_sub = Font(name=F, size=10, italic=True, color="595959")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_input = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="008000")
f_sect = Font(name=F, size=11, bold=True, color="1F3864")
fill_head = PatternFill("solid", fgColor="1F3864")
fill_input = PatternFill("solid", fgColor="FFF2CC")
fill_band = PatternFill("solid", fgColor="F2F2F2")
fill_riise = PatternFill("solid", fgColor="DDEBF7")
fill_riise_head = PatternFill("solid", fgColor="2E75B6")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

LOAD, SCROLL, UPD, OFF, DP, SDK, RAM, RES, INFO = B.DATA_SHEETS
RUN_FIRST, RUN_LAST = 7, 7 + NRUNS - 1  # G..K
COL_MED, COL_MIN, COL_MAX, COL_VAR, COL_NOTE = 12, 13, 14, 15, 16
SDK_COLS = ["App", "Trackers (Exodus)", "Permissions (Exodus)", "Third-party domains on screen (PCAPdroid)",
            "Start-up lock contention (X4)", "Total domains on screen (PCAPdroid)", "Analytics SDKs",
            "Attribution / marketing SDKs", "Crash / APM SDKs", "Other notable SDKs (LibChecker)", "Exodus report URL", "Date checked"]
INFO_COLS = ["App", "Package name", "Version tested", "Date tested", "Framework", "Play Store rating",
             "Play Store downloads", "Download size (MB)", "Installed size (MB)", "Crashes during testing",
             "ANRs during testing", "Top performance complaints in recent reviews"]


def _is_rise(a):
    return B.is_rise(a)


def build_workbook(out, *, apps, device_label, module, metrics=None, weights=None, title=None, subtitle=None):
    """Build the empty workbook for one module and save it. Returns a layout dict for filling cells.

    apps: ordered app names (RISE first). device_label: the phone model. module: loaded module dict.
    metrics: catalogue subset (default: every metric that applies to the module, phone-scoped ones included
    so the _app values can be filled in). weights: {category: weight}.
    """
    weights = weights or B.load_weights()
    if metrics is None:
        metrics = [m for m in B.METRIC_CATALOGUE if B.metric_applies(m, module)[0] or m["phone"]]
    M = [(m["id"], m["category"], B.render(m["name"], module), m["unit"], m["better"], m["scope"], m["tool"],
          B.render(m["how"], module)) for m in metrics]
    devices = [device_label]
    title = title or f"RISE vs Competitors — {module.get('title', module['name'])} Performance Benchmark"
    subtitle = subtitle or (f"Medians from the data tabs, measured on {device_label}. Green = best, red = worst in each row. "
                            "RISE is in the blue column.")
    off_items = [(f"OF{i + 1}", t) for i, t in enumerate(module.get("offline_items", []))]
    dp_items = [(f"DP{i + 1}", module.get("title", module["name"]), t) for i, t in enumerate(module.get("data_items", []))]
    layout = {"runs": {}, "off_rows": {}, "off_cols": {}, "dp_rows": {}, "dp_cols": {}, "sdk_rows": {}, "info_rows": {},
              "off_score_row": None, "dp_score_row": None, "sheets": {}}

    wb = Workbook()

    def setw(ws, widths):
        for col, w in widths.items():
            ws.column_dimensions[col].width = w

    def header(ws, row, cols, start=1):
        for i, h in enumerate(cols):
            c = ws.cell(row=row, column=start + i, value=h)
            c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, box
            if _is_rise(h):
                c.fill = fill_riise_head
        ws.row_dimensions[row].height = 32

    def ttl(ws, t, sub):
        ws["A1"] = t; ws["A1"].font = f_title
        ws["A2"] = sub; ws["A2"].font = f_sub

    def inp(c, numfmt=None):
        c.font, c.fill, c.border = f_input, fill_input, box
        if numfmt:
            c.number_format = numfmt

    def calc(c, numfmt=None):
        c.font, c.border = f_base, box
        if numfmt:
            c.number_format = numfmt

    def default_font(ws):
        for row in ws.iter_rows():
            for c in row:
                if c.font == Font():
                    c.font = f_base

    # ---------------------------------------------------------------- Read Me / Summary placeholders
    rm = wb.active
    rm.title = "Read Me"
    sm = wb.create_sheet("Summary")

    # ---------------------------------------------------------------- Metric Catalogue
    cat = wb.create_sheet("Metric Catalogue")
    ttl(cat, "Metric Catalogue", "Every metric in the benchmark: what it measures, which tool, and exactly how to take the reading.")
    header(cat, 4, ["ID", "Category", "Metric", "Unit", "Better", "Recorded on tab", "Tool", "How to measure"])
    cat_row = {}
    for i, (mid, mcat, name, unit, better, scope, tool, how) in enumerate(M):
        r = 5 + i
        cat_row[mid] = r
        for j, v in enumerate([mid, mcat, name, unit, better, scope[1], tool, how]):
            c = cat.cell(row=r, column=1 + j, value=v)
            c.font, c.border, c.alignment = f_base, box, wrap
    setw(cat, {"A": 6, "B": 20, "C": 44, "D": 7, "E": 8, "F": 20, "G": 20, "H": 90})
    cat.freeze_panes = "A5"

    # ---------------------------------------------------------------- Test Setup
    ts = wb.create_sheet("Test Setup")
    ttl(ts, "Test Setup", "Filled from env.csv by build_report.py. Every result is only comparable if the setup stays fixed.")
    ts["A4"] = "Test device"; ts["A4"].font = f_sect
    header(ts, 5, ["Label", "Model", "Android version", "RAM (GB)", "Render rate (Hz)", "Refresh mode", "Developer", "Notes"])
    for i, d in enumerate(devices):
        r = 6 + i
        c = ts.cell(row=r, column=1, value=d); calc(c); c.font = f_bold
        for col in range(2, 9):
            inp(ts.cell(row=r, column=col))
    ts["A9"] = "Network and conditions"; ts["A9"].font = f_sect
    header(ts, 10, ["Item", "Value", "Notes"])
    for i, item in enumerate(["Wi-Fi RSSI (dBm) / link speed (Mbps)", "Speed test – download (Mbps)", "Private DNS mode",
                              "Battery % / charging", "Thermal status", "Animation scales"]):
        r = 11 + i
        c = ts.cell(row=r, column=1, value=item); calc(c)
        inp(ts.cell(row=r, column=2)); inp(ts.cell(row=r, column=3))
    ts["A17"] = "Test accounts (holdings / watchlist size per app)"; ts["A17"].font = f_sect
    header(ts, 18, ["App", "Version tested", "Logged in with data?", "Holdings / watchlist size", "Notes"])
    for i, a in enumerate(apps):
        r = 19 + i
        c = ts.cell(row=r, column=1, value=a); calc(c); c.font = f_bold
        for col in range(2, 6):
            inp(ts.cell(row=r, column=col))
    layout["setup_app_rows"] = {a: 19 + i for i, a in enumerate(apps)}
    rules_row = 19 + len(apps) + 1
    ts.cell(row=rules_row, column=1, value="Test rules (same for every app)").font = f_sect
    rules = [
        "1. Same phone, same Android version and same Wi-Fi for every app in this module. Don't update the OS mid-benchmark.",
        "2. Before each session: battery above 50%, brightness fixed, all other apps closed, Do Not Disturb on, Power saving off.",
        "3. Five runs per timing metric, three for live sessions, one for facts. The sheet uses the MEDIAN.",
        "4. If a row's Variation (max−min ÷ median) turns red (above 20%), check for interference and redo that row.",
        "5. Live-data tests (U1–U5, C1–C4) only during market hours (09:15–15:30 IST).",
        "6. Phone-scoped metrics (L1, L2, R1, R4–R8, X1, X2, X4, A3, A4) come from the _app module, measured once per phone.",
        "7. Keep adb logcat -b crash running in every session to catch crashes and ANRs.",
    ]
    for i, t in enumerate(rules):
        ts.cell(row=rules_row + 1 + i, column=1, value=t).font = f_base
    setw(ts, {"A": 34, "B": 24, "C": 24, "D": 24, "E": 24, "F": 22, "G": 22, "H": 30})

    # ---------------------------------------------------------------- runs sheets
    runs_ranges = {}

    def runs_sheet(name, sub):
        ws = wb.create_sheet(name)
        ttl(ws, name, sub)
        header(ws, 4, ["ID", "Metric", "Unit", "App", "Device", "Lookup key"] +
               [f"Run {i}" for i in range(1, NRUNS + 1)] + ["Median", "Min", "Max", "Variation", "Notes"])
        mets = [m for m in M if m[5] == ("runs", name)]
        r = 5
        num_dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                error="Enter a number (no units).", showErrorMessage=True)
        ws.add_data_validation(num_dv)
        for bi, (mid, _, mname, unit, *_rest) in enumerate(mets):
            band = fill_band if bi % 2 else None
            for a in apps:
                for d in devices:
                    for j, v in enumerate([mid, mname, unit, a, d]):
                        c = ws.cell(row=r, column=1 + j, value=v); calc(c)
                        if band: c.fill = band
                        if _is_rise(a): c.fill = fill_riise
                    c = ws.cell(row=r, column=6, value=f'=A{r}&"|"&D{r}&"|"&E{r}')
                    calc(c); c.font = Font(name=F, size=8, color="808080")
                    for col in range(RUN_FIRST, RUN_LAST + 1):
                        inp(ws.cell(row=r, column=col), "#,##0.0")
                        num_dv.add(ws.cell(row=r, column=col))
                    rng = f"{L(RUN_FIRST)}{r}:{L(RUN_LAST)}{r}"
                    calc(ws.cell(row=r, column=COL_MED, value=f'=IF(COUNT({rng})=0,"",MEDIAN({rng}))'), "#,##0.0")
                    calc(ws.cell(row=r, column=COL_MIN, value=f'=IF(COUNT({rng})=0,"",MIN({rng}))'), "#,##0.0")
                    calc(ws.cell(row=r, column=COL_MAX, value=f'=IF(COUNT({rng})=0,"",MAX({rng}))'), "#,##0.0")
                    calc(ws.cell(row=r, column=COL_VAR,
                                 value=f'=IF(COUNT({rng})<2,"",IF(L{r}=0,"",(N{r}-M{r})/L{r}))'), "0%")
                    ws.cell(row=r, column=COL_MED).font = f_bold
                    inp(ws.cell(row=r, column=COL_NOTE))
                    layout["runs"][(mid, a)] = (name, r)
                    r += 1
        last = r - 1
        if last < 5:  # every metric on this sheet is skipped by the module: leave the header only
            ws.cell(row=5, column=2, value="No metrics on this sheet for this module (skipped in the module file).")
            last = 5
        ws.conditional_formatting.add(f"O5:O{last}", CellIsRule(operator="greaterThan", formula=["0.2"],
                                      fill=PatternFill("solid", fgColor="F8CBAD"), font=Font(name=F, color="9C0006")))
        setw(ws, {"A": 6, "B": 46, "C": 7, "D": 13, "E": 22, "F": 14, "G": 9, "H": 9, "I": 9, "J": 9, "K": 9,
                  "L": 10, "M": 9, "N": 9, "O": 10, "P": 36})
        ws.freeze_panes = "G5"
        ws.auto_filter.ref = f"A4:P{last}"
        runs_ranges[name] = max(last, 5)
        return ws

    screen = module.get("rise", {}).get("screen", "the screen")
    lst = module.get("rise", {}).get("list", "the list")
    runs_sheet(LOAD, f"Loading time for {screen}, API/network behaviour, user-wise (account size) loads, offline recovery. Filled from runs.csv; Median/Min/Max/Variation are formulas.")
    runs_sheet(SCROLL, f"Scrolling smoothness on {lst} and gesture jank. Perfetto FrameTimeline for every app (bench.py).")
    runs_sheet(UPD, "Live screen updation during market hours (09:15–15:30 IST), idle redraw rate and wasted redraws.")

    # ---------------------------------------------------------------- 4 Offline
    off = wb.create_sheet(OFF, index=wb.sheetnames.index(UPD) + 1)
    ttl(off, OFF, f"Caching behaviour in offline mode on {screen}. Yes / Partial / No / N/A per item, from record.py; the score is a formula.")
    header(off, 4, ["ID", "Scenario"] + apps)
    ypn = DataValidation(type="list", formula1='"Yes,Partial,No,N/A"', allow_blank=True)
    off.add_data_validation(ypn)
    for i, (sid, s) in enumerate(off_items):
        r = 5 + i
        calc(off.cell(row=r, column=1, value=sid)); c = off.cell(row=r, column=2, value=s); calc(c); c.alignment = wrap
        layout["off_rows"][B.slug(s)] = r
        for j in range(len(apps)):
            c = off.cell(row=r, column=3 + j); inp(c); c.alignment = center; ypn.add(c)
    for j, a in enumerate(apps):
        layout["off_cols"][a] = 3 + j
    off_last = max(5 + len(off_items) - 1, 5)
    OFF_SCORE_ROW = off_last + 1
    layout["off_score_row"] = OFF_SCORE_ROW
    c = off.cell(row=OFF_SCORE_ROW, column=2, value="Offline caching score (Yes=2, Partial=1, No=0; N/A excluded)"); calc(c); c.font = f_bold
    for j in range(len(apps)):
        col = L(3 + j); rng = f"{col}5:{col}{off_last}"
        y, p, n = (f'COUNTIF({rng},"Yes")', f'COUNTIF({rng},"Partial")', f'COUNTIF({rng},"No")')
        c = off.cell(row=OFF_SCORE_ROW, column=3 + j, value=f'=IF({y}+{p}+{n}=0,"",(2*{y}+{p})/(2*({y}+{p}+{n})))')
        calc(c, "0%"); c.font = f_bold; c.alignment = center
    grid_top = OFF_SCORE_ROW + 3
    off.cell(row=grid_top - 1, column=1, value="Observations (notes recorded with each item)").font = f_sect
    header(off, grid_top, ["ID", "Scenario"] + apps)
    for i, (sid, s) in enumerate(off_items):
        r = grid_top + 1 + i
        calc(off.cell(row=r, column=1, value=sid)); c = off.cell(row=r, column=2, value=f"=B{5 + i}"); calc(c); c.alignment = wrap
        for j in range(len(apps)):
            c = off.cell(row=r, column=3 + j); inp(c); c.alignment = wrap
    layout["off_note_rows"] = {B.slug(s): grid_top + 1 + i for i, (sid, s) in enumerate(off_items)}
    last_app_col = L(2 + len(apps))
    rr = f"C5:{last_app_col}{off_last}"
    off.conditional_formatting.add(rr, CellIsRule(operator="equal", formula=['"Yes"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    off.conditional_formatting.add(rr, CellIsRule(operator="equal", formula=['"Partial"'], fill=PatternFill("solid", fgColor="FFEB9C")))
    off.conditional_formatting.add(rr, CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid", fgColor="FFC7CE")))
    setw(off, {"A": 6, "B": 58, **{L(3 + j): 16 for j in range(len(apps))}})
    off.freeze_panes = "C5"

    # ---------------------------------------------------------------- 5 Data Pointers
    dp = wb.create_sheet(DP, index=wb.sheetnames.index(OFF) + 1)
    ttl(dp, DP, f"Which data points and features each app offers on {screen}. Yes / Partial / No / N/A per item, from record.py.")
    header(dp, 4, ["ID", "Group", "Data pointer / feature"] + apps)
    for i, (did, grp, it) in enumerate(dp_items):
        r = 5 + i
        calc(dp.cell(row=r, column=1, value=did)); calc(dp.cell(row=r, column=2, value=grp))
        c = dp.cell(row=r, column=3, value=it); calc(c)
        layout["dp_rows"][B.slug(it)] = r
        for j in range(len(apps)):
            c = dp.cell(row=r, column=4 + j); inp(c); c.alignment = center
    for j, a in enumerate(apps):
        layout["dp_cols"][a] = 4 + j
    dp_last = max(5 + len(dp_items) - 1, 5)
    dp_last_col = L(3 + len(apps))
    dp_dv = DataValidation(type="list", formula1='"Yes,Partial,No,N/A"', allow_blank=True)
    dp.add_data_validation(dp_dv); dp_dv.add(f"D5:{dp_last_col}{dp_last}")
    DP_SCORE_ROW = dp_last + 1
    layout["dp_score_row"] = DP_SCORE_ROW
    c = dp.cell(row=DP_SCORE_ROW, column=3, value="Coverage (Yes=1, Partial=0.5, No=0; N/A excluded)"); calc(c); c.font = f_bold
    for j in range(len(apps)):
        col = L(4 + j); rng = f"{col}5:{col}{dp_last}"
        y, p, nn = (f'COUNTIF({rng},"Yes")', f'COUNTIF({rng},"Partial")', f'COUNTIF({rng},"No")')
        c = dp.cell(row=DP_SCORE_ROW, column=4 + j, value=f'=IF({y}+{p}+{nn}=0,"",({y}+0.5*{p})/({y}+{p}+{nn}))')
        calc(c, "0%"); c.font = f_bold; c.alignment = center
    rr = f"D5:{dp_last_col}{dp_last}"
    dp.conditional_formatting.add(rr, CellIsRule(operator="equal", formula=['"Yes"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    dp.conditional_formatting.add(rr, CellIsRule(operator="equal", formula=['"Partial"'], fill=PatternFill("solid", fgColor="FFEB9C")))
    dp.conditional_formatting.add(rr, CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid", fgColor="FFC7CE")))
    setw(dp, {"A": 6, "B": 18, "C": 44, **{L(4 + j): 14 for j in range(len(apps))}})
    dp.freeze_panes = "D5"

    # ---------------------------------------------------------------- 6 SDKs
    sdk = wb.create_sheet(SDK, index=wb.sheetnames.index(DP) + 1)
    ttl(sdk, SDK, "External SDK footprint. Counts from Exodus Privacy, PCAPdroid and bench.py (X4); SDK names from Exodus and LibChecker.")
    header(sdk, 4, SDK_COLS)
    for i, a in enumerate(apps):
        r = 5 + i
        c = sdk.cell(row=r, column=1, value=a); calc(c); c.font = f_bold
        if _is_rise(a): c.fill = fill_riise
        for col in range(2, len(SDK_COLS) + 1):
            c = sdk.cell(row=r, column=col); inp(c, "#,##0" if col <= 6 else ("dd-mmm-yyyy" if col == len(SDK_COLS) else None))
            c.alignment = wrap
        layout["sdk_rows"][a] = r
    setw(sdk, {"A": 14, "B": 12, "C": 13, "D": 18, "E": 16, "F": 16, "G": 26, "H": 26, "I": 22, "J": 30, "K": 30, "L": 13})
    sdk.freeze_panes = "B5"

    runs_sheet(RAM, "RAM (Total PSS) with the Java/native/code/graphics split in Notes, the leak check, and storage. Storage rows need Run 1 only.")
    runs_sheet(RES, f"CPU, battery, temperature and data use during a live session on {screen} (market hours).")

    # ---------------------------------------------------------------- 9 App Info & Stability
    info = wb.create_sheet(INFO)
    ttl(info, INFO, "App version under test, framework, Play Store data and stability during testing.")
    header(info, 4, INFO_COLS)
    fw = DataValidation(type="list", formula1='"Native (Kotlin/Java),Flutter,React Native,Hybrid / WebView,Unknown"', allow_blank=True)
    info.add_data_validation(fw)
    fmts = {4: "dd-mmm-yyyy", 6: "0.0", 8: "#,##0.0", 9: "#,##0.0", 10: "0", 11: "0"}
    for i, a in enumerate(apps):
        r = 5 + i
        c = info.cell(row=r, column=1, value=a); calc(c); c.font = f_bold
        if _is_rise(a): c.fill = fill_riise
        for col in range(2, 13):
            c = info.cell(row=r, column=col); inp(c, fmts.get(col)); c.alignment = wrap
        fw.add(info.cell(row=r, column=5))
        layout["info_rows"][a] = r
    note_row = 5 + len(apps) + 1
    info.cell(row=note_row, column=1, value="Perfetto FrameTimeline is used for the FPS metrics of every app. LibChecker shows the framework.").font = f_sub
    setw(info, {"A": 14, "B": 26, "C": 13, "D": 13, "E": 20, "F": 11, "G": 13, "H": 12, "I": 12, "J": 12, "K": 12, "L": 44})
    info.freeze_panes = "B5"
    n_app_rows = len(apps)
    app_last = 4 + n_app_rows

    # ---------------------------------------------------------------- Summary
    ttl(sm, title, subtitle)
    sm["A3"] = "Device:"; sm["A3"].font = f_bold
    sm["C3"] = device_label; calc(sm["C3"]); sm["C3"].font = Font(name=F, size=11, bold=True, color="1F3864")
    sm["D3"] = "← the phone this module was measured on (one per module; see README on comparing across phones)"; sm["D3"].font = f_sub
    HR = 6
    cols = ["ID", "Category", "Metric", "Unit", "Better"] + apps + ["Best app", "RISE rank", "Apps measured", "RISE gap vs best"]
    header(sm, HR, cols)
    first = HR + 1
    APP0, APPN = 6, 6 + len(apps) - 1
    c_best, c_rank, c_cnt, c_gap = APPN + 1, APPN + 2, APPN + 3, APPN + 4
    fa, la = L(APP0), L(APPN)
    rise_col = L(APP0 + next((k for k, a in enumerate(apps) if _is_rise(a)), 0))
    for i, (mid, mcat, name, unit, better, scope, *_r) in enumerate(M):
        r = first + i
        cr = cat_row[mid]
        for j, colname in enumerate("ABCDE"):
            c = sm.cell(row=r, column=1 + j, value=f"='Metric Catalogue'!{colname}{cr}")
            calc(c); c.font = f_link
        sm.cell(row=r, column=3).alignment = wrap
        for k, a in enumerate(apps):
            col = APP0 + k
            hdr = f"{L(col)}${HR}"
            if scope[0] == "runs":
                sh, last = scope[1], runs_ranges[scope[1]]
                f = (f"=IFERROR(INDEX('{sh}'!$L$5:$L${last},MATCH($A{r}&\"|\"&{hdr}&\"|\"&$C$3,"
                     f"'{sh}'!$F$5:$F${last},0)),\"\")")
            elif scope[0] == "app":
                sh, colx = scope[1], scope[2]
                look = f"INDEX('{sh}'!${colx}$5:${colx}${app_last},MATCH({hdr},'{sh}'!$A$5:$A${app_last},0))"
                f = f'=IFERROR(IF({look}="","",{look}),"")'
            else:
                sh = scope[1]
                row_s, c0, c1 = (OFF_SCORE_ROW, "C", last_app_col) if sh == OFF else (DP_SCORE_ROW, "D", dp_last_col)
                f = (f"=IFERROR(INDEX('{sh}'!${c0}${row_s}:${c1}${row_s},MATCH({hdr},"
                     f"'{sh}'!${c0}$4:${c1}$4,0)),\"\")")
            c = sm.cell(row=r, column=col, value=f)
            calc(c, "0%" if unit == "%" and scope[0] == "score" else "#,##0.0")
            c.alignment = Alignment(horizontal="center")
            c.font = Font(name=F, size=10, bold=True, color="008000") if _is_rise(a) else f_link
        rng = f"{fa}{r}:{la}{r}"
        best = f'IF($E{r}="Lower",MIN({rng}),MAX({rng}))'
        calc(sm.cell(row=r, column=c_best,
                     value=f'=IF(COUNT({rng})=0,"",INDEX(${fa}${HR}:${la}${HR},MATCH({best},{rng},0)))'))
        calc(sm.cell(row=r, column=c_rank,
                     value=f'=IF(OR({rise_col}{r}="",COUNT({rng})=0),"",IF($E{r}="Lower",COUNTIF({rng},"<"&{rise_col}{r}),COUNTIF({rng},">"&{rise_col}{r}))+1)'))
        calc(sm.cell(row=r, column=c_cnt, value=f"=COUNT({rng})"))
        calc(sm.cell(row=r, column=c_gap,
                     value=f'=IF(OR({rise_col}{r}="",COUNT({rng})<2),"",IF($E{r}="Lower",IF(MIN({rng})=0,"",({rise_col}{r}-MIN({rng}))/MIN({rng})),IF(MAX({rng})=0,"",(MAX({rng})-{rise_col}{r})/MAX({rng}))))'),
             "0%")
        for cc in (c_best, c_rank, c_cnt, c_gap):
            sm.cell(row=r, column=cc).alignment = Alignment(horizontal="center")
        sm.cell(row=r, column=c_best).font = f_bold
        green, yellow, red = "63BE7B", "FFEB84", "F8696B"
        lo, hi = (green, red) if better == "Lower" else (red, green)
        sm.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color=lo, mid_type="percentile",
                                                           mid_value=50, mid_color=yellow, end_type="max", end_color=hi))
        if i % 2:
            for cc in range(1, 6):
                sm.cell(row=r, column=cc).fill = fill_band
    last_sm = first + len(M) - 1
    sm.conditional_formatting.add(f"{L(c_best)}{first}:{L(c_best)}{last_sm}",
                                  CellIsRule(operator="equal", formula=['"RISE"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    sm.conditional_formatting.add(f"{L(c_gap)}{first}:{L(c_gap)}{last_sm}",
                                  CellIsRule(operator="greaterThan", formula=["0.2"], font=Font(name=F, color="C00000", bold=True)))
    sm["A4"] = "RISE is best in:"; sm["A4"].font = f_bold
    bc = L(c_best)
    sm["C4"] = f'=COUNTIF({bc}{first}:{bc}{last_sm},"RISE")&" of "&COUNTIF({L(c_cnt)}{first}:{L(c_cnt)}{last_sm},">0")&" metrics measured"'
    calc(sm["C4"]); sm["C4"].font = f_bold; sm["C4"].border = Border()
    sm["H4"] = "Average RISE rank:"; sm["H4"].font = f_bold
    rc = L(c_rank)
    sm["J4"] = f'=IFERROR(AVERAGE({rc}{first}:{rc}{last_sm}),"–")'
    calc(sm["J4"], "0.0"); sm["J4"].font = f_bold; sm["J4"].border = Border()
    sm[f"A{last_sm + 2}"] = ("RISE rank: 1 = best of the apps measured. RISE gap vs best: how far RISE is behind the best app "
                             "(0% = RISE is best; above 20% shown in red). Unranked apps (RISE UAT) are shown for reference only.")
    sm[f"A{last_sm + 2}"].font = f_sub
    setw(sm, {"A": 6, "B": 18, "C": 44, "D": 7, "E": 8, **{L(c): 12 for c in range(APP0, APPN + 1)},
              L(c_best): 13, L(c_rank): 10, L(c_cnt): 10, L(c_gap): 12})
    sm.freeze_panes = sm.cell(row=first, column=APP0)
    sm.auto_filter.ref = f"A{HR}:{L(c_gap)}{last_sm}"

    # ---------------------------------------------------------------- Read Me content
    ttl(rm, title, f"Module {module.get('title', module['name'])} on {device_label}. Generated by build_report.py from runs.csv; do not edit by hand.")
    rm["A4"] = "How to read this workbook"; rm["A4"].font = f_sect
    steps = [
        "1. 'Test Setup' shows the phone, network and accounts, from env.csv.",
        "2. 'Metric Catalogue' lists every metric, the tool, and the exact method (same text as METRICS.md).",
        "3. Data tabs (1–9) hold every run next to its median, so any figure can be traced to runs.csv and its artifact.",
        "4. 'Summary' shows medians per app, the best app, RISE's rank and gap per metric.",
        "5. 'Scorecard' turns everything into one 0–100 score per app and ranks them. Weights come from weights.json.",
        "6. Blank cells are metrics not yet measured. Metrics that do not apply to this module are left out (see METRICS.md).",
    ]
    for i, s in enumerate(steps):
        rm.cell(row=5 + i, column=1, value=s).font = f_base
    rm["A12"] = "Legend"; rm["A12"].font = f_sect
    c = rm["A13"]; c.value = "Yellow cell, blue text"; inp(c); rm["B13"] = "Value from runs.csv / env.csv"
    c = rm["A14"]; c.value = "Black text"; calc(c); rm["B14"] = "Formula"
    c = rm["A15"]; c.value = "Green text"; calc(c); c.font = f_link; rm["B15"] = "Pulled from another tab"
    c = rm["A16"]; c.value = "Blue rows / column"; calc(c); c.fill = fill_riise; rm["B16"] = "RISE (our app)"
    c = rm["A17"]; c.value = "Red Variation"; calc(c); c.fill = PatternFill("solid", fgColor="F8CBAD"); rm["B17"] = "Runs differ by more than 20%: redo that row"
    for r in range(13, 18):
        rm.cell(row=r, column=2).font = f_base
    rm["A19"] = "Tools"; rm["A19"].font = f_sect
    tools = [
        ("ADB (Android Platform Tools)", "Mac", "Launch time, lock contention, RAM, bytes, CPU, battery, crashes", "brew install android-platform-tools"),
        ("scrcpy", "Mac", "Screen recording to count frames for screen-load and tick timings", "brew install scrcpy"),
        ("ffmpeg", "Mac", "Wasted-redraw analysis of idle recordings (U5)", "brew install ffmpeg"),
        ("Perfetto (Python package)", "Mac", "Reads the frame traces bench.py records", "pip3 install --user perfetto"),
        ("bench.py / record.py / preflight.py", "Mac", "Automates the measured metrics and records the manual ones", "In this folder"),
        ("PCAPdroid", "Phone", "Connections and domains per app, without root", "Play Store"),
        ("LibChecker", "Phone", "SDKs, native libraries, framework", "Play Store / F-Droid"),
        ("Exodus Privacy", "Web", "Trackers and permissions in each Play Store APK", "exodus-privacy.eu.org"),
    ]
    header(rm, 20, ["Tool", "Runs on", "Used for", "Get it"])
    for i, row in enumerate(tools):
        for j, v in enumerate(row):
            c = rm.cell(row=21 + i, column=1 + j, value=v); calc(c); c.alignment = wrap
    setw(rm, {"A": 34, "B": 40, "C": 44, "D": 34})

    # ---------------------------------------------------------------- Scorecard
    sc = wb.create_sheet("Scorecard", wb.sheetnames.index("Summary") + 1)
    ttl(sc, "Scorecard — Overall Ranking", "Each metric is scored 0–100 (best app = 100, worst = 0), averaged per category, then weighted into one overall score. Weights come from weights.json.")
    sc["A3"] = f'="Device: "&Summary!C3'
    sc["A3"].font = f_bold
    CATS = []
    for m in M:
        if m[1] not in CATS:
            CATS.append(m[1])
    NA = len(apps)
    AC0 = 3
    acol = lambda k: L(AC0 + k)
    fa_sc, la_sc = acol(0), acol(NA - 1)
    LB_H = 5
    sc.cell(row=LB_H - 1, column=1, value="Leaderboard").font = f_sect
    header(sc, LB_H, ["Rank", "App", "Overall score (0–100)", "Data coverage"])
    CT_H = LB_H + NA + 4
    sc.cell(row=CT_H - 1, column=1, value="Category scores (weights from weights.json, read-only here)").font = f_sect
    header(sc, CT_H, ["Category", "Weight"] + apps)
    ct_first, ct_last = CT_H + 1, CT_H + len(CATS)
    TOT, OVR, COV, RNK, ADJ = ct_last + 1, ct_last + 2, ct_last + 3, ct_last + 4, ct_last + 5
    MG_H = ADJ + 4
    sc.cell(row=MG_H - 1, column=1, value="Metric scores (0–100, from the Summary tab)").font = f_sect
    header(sc, MG_H, ["ID", "Category", "Metric"] + apps)
    mg_first, mg_last = MG_H + 1, MG_H + len(M)
    HC0 = AC0 + NA + 2
    HU0 = HC0 + NA + 1
    sc.cell(row=CT_H - 1, column=HC0, value="Helper: weight × score (don't edit)").font = f_sub
    sc.cell(row=CT_H - 1, column=HU0, value="Helper: weight used (don't edit)").font = f_sub
    for k, a in enumerate(apps):
        for base in (HC0, HU0):
            c = sc.cell(row=CT_H, column=base + k, value=a); c.font = Font(name=F, size=8, bold=True, color="808080")
    for i, m in enumerate(M):
        r, s_row = mg_first + i, first + i
        for j, colname in enumerate("ABC"):
            c = sc.cell(row=r, column=1 + j, value=f"=Summary!{colname}{s_row}"); calc(c); c.font = f_link
        rng = f"Summary!${fa}{s_row}:${la}{s_row}"
        mx, mn = f"MAX({rng})", f"MIN({rng})"
        for k in range(NA):
            x = f"Summary!{L(APP0 + k)}{s_row}"
            f = (f'=IF(OR({x}="",Summary!${L(c_cnt)}{s_row}<2),"",IF({mx}={mn},100,'
                 f'IF(Summary!$E{s_row}="Lower",({mx}-{x})/({mx}-{mn}),({x}-{mn})/({mx}-{mn}))*100))')
            c = sc.cell(row=r, column=AC0 + k, value=f); calc(c, "0"); c.alignment = Alignment(horizontal="center")
    for i, catn in enumerate(CATS):
        r = ct_first + i
        c = sc.cell(row=r, column=1, value=catn); calc(c); c.font = f_bold
        calc(sc.cell(row=r, column=2, value=weights.get(catn, 0)), "0")
        for k in range(NA):
            col = acol(k)
            c = sc.cell(row=r, column=AC0 + k,
                        value=f'=IFERROR(AVERAGEIFS({col}${mg_first}:{col}${mg_last},$B${mg_first}:$B${mg_last},$A{r}),"")')
            calc(c, "0"); c.alignment = Alignment(horizontal="center")
            h = sc.cell(row=r, column=HC0 + k, value=f'=IF(ISNUMBER({col}{r}),$B{r}*{col}{r},"")')
            h.font = Font(name=F, size=8, color="808080")
            u = sc.cell(row=r, column=HU0 + k, value=f'=IF(ISNUMBER({col}{r}),$B{r},"")')
            u.font = Font(name=F, size=8, color="808080")
    c = sc.cell(row=TOT, column=1, value="Total weight"); calc(c); c.font = f_bold
    c = sc.cell(row=TOT, column=2, value=f"=SUM(B{ct_first}:B{ct_last})"); calc(c, "0"); c.font = f_bold
    for row_, label in ((OVR, "Overall score"), (COV, "Data coverage (share of weight measured)"), (RNK, "Rank")):
        c = sc.cell(row=row_, column=1, value=label); calc(c); c.font = f_bold
    for k in range(NA):
        hc, hu, col = L(HC0 + k), L(HU0 + k), acol(k)
        c = sc.cell(row=OVR, column=AC0 + k,
                    value=f'=IF(SUM({hu}{ct_first}:{hu}{ct_last})=0,"",SUM({hc}{ct_first}:{hc}{ct_last})/SUM({hu}{ct_first}:{hu}{ct_last}))')
        calc(c, "0.0"); c.font = f_bold; c.alignment = Alignment(horizontal="center")
        c = sc.cell(row=COV, column=AC0 + k,
                    value=f'=IF($B${TOT}=0,"",SUM({hu}{ct_first}:{hu}{ct_last})/$B${TOT})')
        calc(c, "0%"); c.alignment = Alignment(horizontal="center")
        c = sc.cell(row=RNK, column=AC0 + k,
                    value=f'=IF(ISNUMBER({col}{OVR}),COUNTIF(${fa_sc}${OVR}:${la_sc}${OVR},">"&{col}{OVR})+1,"")')
        calc(c, "0"); c.font = f_bold; c.alignment = Alignment(horizontal="center")
        c = sc.cell(row=ADJ, column=AC0 + k, value=f'=IF(ISNUMBER({col}{OVR}),{col}{OVR}+{NA - k}/1000000,"")')
        c.font = Font(name=F, size=8, color="808080"); c.number_format = "0.0"
    sc.cell(row=ADJ, column=1, value="Helper: tie-break (don't edit)").font = Font(name=F, size=8, color="808080")
    adj_rng = f"${fa_sc}${ADJ}:${la_sc}${ADJ}"
    for k in range(NA):
        r = LB_H + 1 + k
        c = sc.cell(row=r, column=1, value=k + 1); calc(c); c.alignment = Alignment(horizontal="center")
        big = f"LARGE({adj_rng},{k + 1})"
        pos = f"MATCH({big},{adj_rng},0)"
        c = sc.cell(row=r, column=2, value=f'=IFERROR(INDEX(${fa_sc}${CT_H}:${la_sc}${CT_H},{pos}),"")'); calc(c); c.font = f_bold
        c = sc.cell(row=r, column=3, value=f'=IFERROR(INDEX(${fa_sc}${OVR}:${la_sc}${OVR},{pos}),"")'); calc(c, "0.0")
        c.alignment = Alignment(horizontal="center")
        c = sc.cell(row=r, column=4, value=f'=IFERROR(INDEX(${fa_sc}${COV}:${la_sc}${COV},{pos}),"")'); calc(c, "0%")
        c.alignment = Alignment(horizontal="center")
    lb_first, lb_last = LB_H + 1, LB_H + NA
    sc.conditional_formatting.add(f"B{lb_first}:D{lb_last}", FormulaRule(formula=[f'$B{lb_first}="RISE"'], fill=fill_riise, font=Font(name=F, bold=True, color="1F3864")))
    scale = ColorScaleRule(start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=50,
                           mid_color="FFEB84", end_type="num", end_value=100, end_color="63BE7B")
    for rng_ in (f"C{lb_first}:C{lb_last}", f"{fa_sc}{ct_first}:{la_sc}{ct_last}", f"{fa_sc}{OVR}:{la_sc}{OVR}",
                 f"{fa_sc}{mg_first}:{la_sc}{mg_last}"):
        sc.conditional_formatting.add(rng_, scale)
    sc.conditional_formatting.add(f"B{TOT}", CellIsRule(operator="notEqual", formula=["100"],
                                  fill=PatternFill("solid", fgColor="F8CBAD"), font=Font(name=F, bold=True, color="9C0006")))
    sc.conditional_formatting.add(f"{fa_sc}{RNK}:{la_sc}{RNK}", CellIsRule(operator="equal", formula=["1"],
                                  fill=PatternFill("solid", fgColor="C6EFCE")))
    for k, a in enumerate(apps):
        if _is_rise(a):
            for r in list(range(ct_first, ct_last + 1)) + [OVR, COV, RNK] + list(range(mg_first, mg_last + 1)):
                sc.cell(row=r, column=AC0 + k).fill = fill_riise
    ch = BarChart()
    ch.type = "bar"
    ch.title = "Overall score (0–100)"
    ch.legend = None
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 100
    ch.x_axis.scaling.orientation = "maxMin"
    ch.add_data(Reference(sc, min_col=3, min_row=lb_first, max_row=lb_last), titles_from_data=False)
    ch.set_categories(Reference(sc, min_col=2, min_row=lb_first, max_row=lb_last))
    ch.height, ch.width = 7, 14
    sc.add_chart(ch, "F4")
    sc[f"A{mg_last + 2}"] = ("Method: for each metric, score = 100 × (worst − app) ÷ (worst − best), so the best app gets 100 and the "
                            "worst gets 0 (needs at least 2 apps measured). Category score = average of its metric scores. Overall = "
                            "weighted average of category scores, using only categories with data. 'Data coverage' shows how much "
                            "of the total weight was measured for that app; compare apps fairly only when coverage is similar. "
                            "Unranked apps (RISE UAT) take part in these formulas only if present; the combined report excludes them from ranks.")
    sc[f"A{mg_last + 2}"].font = f_sub
    setw(sc, {"A": 26, "B": 18, "C": 44, **{acol(k): 12 for k in range(NA)}})
    sc.column_dimensions["C"].width = 22
    sc.freeze_panes = "A5"

    for ws in wb.worksheets:
        default_font(ws)
        ws.sheet_view.showGridLines = False
    rm.sheet_properties.tabColor = "808080"
    sm.sheet_properties.tabColor = "2E75B6"
    sc.sheet_properties.tabColor = "2E75B6"
    for n in B.DATA_SHEETS:
        wb[n].sheet_properties.tabColor = "FFC000"
    wb.active = wb.sheetnames.index("Summary")
    wb.save(out)
    layout["metrics"] = [m[0] for m in M]
    return layout


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("usage: python3 build_benchmark.py out.xlsx [module]")
    mod = B.load_module(sys.argv[2] if len(sys.argv) > 2 else "portfolio")
    apps = [a for a in B.load_apps() if B.load_apps()[a].get("ranked", True)]
    build_workbook(sys.argv[1], apps=apps, device_label="Phone", module=mod)
    print("saved", sys.argv[1])
