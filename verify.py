#!/usr/bin/env python3
"""RISE re-check: measure only RISE again on the report's phone, compare it with an existing module
report, write the new RISE values into that same workbook, and say what improved.

  python3 verify.py baseline --report reports/quote/quote_<dev>_<device>_<date>.xlsx [--module quote]
        RISE values in the report (median per metric), its phone and RISE version, and the bench.py
        groups to run again. Nothing is written.

  python3 verify.py compare --report <xlsx> --module quote [--date 2026-10-23] [--update]
        Reads the new RISE runs from results/_verify/<date>/<module>/<developer>_<device>/ (and the
        _app folder beside it), compares per metric, category score, overall score and rank, and writes
        reports/<module>/verify_<module>_<developer>_<device>_<date>.md.
        --update also writes the new RISE values into the same workbook (a backup is kept next to it)
        and adds a 'RISE Verification' sheet. Competitor values are never changed.

New RISE runs go to their own folder, so the results behind the report stay as they were:
  python3 bench.py run RISE --module quote --results-dir results/_verify/<date>
  python3 bench.py run RISE --module _app --results-dir results/_verify/<date>
  python3 record.py --module quote --app RISE L14 --value 1840 --results-dir results/_verify/<date>

The re-check must run on the same phone model as the report: raw values from different phones are
not comparable, so compare stops on a mismatch and --update is refused.
"""
import argparse
import datetime
import shutil
import statistics
import sys
from pathlib import Path

import benchlib as B
import build_report as R

VERIFY_DIR = B.RESULTS_DIR / "_verify"
NOISE_PCT = 5.0          # a median that moved less than this (in %) counts as no change
HEAD_ROW, FIRST_ROW = 4, 5   # header and first data row of every data tab
VERIFY_SHEET = "RISE Verification"
RISE = "RISE"


# ---------------------------------------------------------------- reading the report (pure on a workbook)

def _num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def parse_workbook(wb):
    """Everything a module workbook holds per app, with cell positions for writing back.

    {"device", "apps", "runs": {(mid, app): {"sheet", "row", "runs": [5 values]}},
     "facts": {(mid, app): {"sheet", "cell", "value"}}, "items": {"O1"|"D1": {slug: {app: {"sheet", "cell", "value"}}}},
     "info": {app: {"version", "date", "row"}}}
    """
    rep = {"device": wb["Summary"]["C3"].value if "Summary" in wb.sheetnames else None,
           "apps": [], "runs": {}, "facts": {}, "items": {"O1": {}, "D1": {}}, "info": {}}
    run_sheets = {m["scope"][1] for m in B.METRIC_CATALOGUE if m["scope"][0] == "runs"}
    for name in run_sheets & set(wb.sheetnames):
        ws = wb[name]
        for r in range(FIRST_ROW, ws.max_row + 1):
            mid, app = ws.cell(row=r, column=1).value, ws.cell(row=r, column=4).value
            if mid in B.METRICS and app:
                runs = [ws.cell(row=r, column=R.RUN_FIRST + k).value for k in range(B.NRUNS)]
                rep["runs"][(mid, B.canonical_app(app))] = {"sheet": name, "row": r, "runs": runs}
    for m in B.METRIC_CATALOGUE:
        if m["scope"][0] != "app" or m["scope"][1] not in wb.sheetnames:
            continue
        ws, col = wb[m["scope"][1]], m["scope"][2]
        for r in range(FIRST_ROW, ws.max_row + 1):
            app = ws.cell(row=r, column=1).value
            if app:
                rep["facts"][(m["id"], B.canonical_app(app))] = {"sheet": ws.title, "cell": f"{col}{r}",
                                                               "value": ws[f"{col}{r}"].value}
    for base, sheet, text_col, prefix in (("O1", B.OFF, 2, "OF"), ("D1", B.DP, 3, "DP")):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        heads = [ws.cell(row=HEAD_ROW, column=c).value for c in range(1, ws.max_column + 1)]
        app_cols = {B.canonical_app(h): c + 1 for c, h in enumerate(heads) if c + 1 > text_col and h}
        for r in range(FIRST_ROW, ws.max_row + 1):
            rid = ws.cell(row=r, column=1).value
            if not (isinstance(rid, str) and rid.startswith(prefix)):
                break  # the score row ends the item block; the observations block below repeats the IDs
            s = B.slug(ws.cell(row=r, column=text_col).value or "")
            rep["items"][base][s] = {app: {"sheet": sheet, "cell": ws.cell(row=r, column=c).coordinate,
                                           "value": ws.cell(row=r, column=c).value} for app, c in app_cols.items()}
    if B.INFO in wb.sheetnames:
        ws = wb[B.INFO]
        for r in range(FIRST_ROW, ws.max_row + 1):
            app = ws.cell(row=r, column=1).value
            if app:
                rep["info"][B.canonical_app(app)] = {"row": r, "version": ws.cell(row=r, column=3).value,
                                                     "date": ws.cell(row=r, column=4).value}
    if "Summary" in wb.sheetnames:
        ws = wb["Summary"]
        for c in range(6, ws.max_column + 1):
            h = ws.cell(row=6, column=c).value
            if not h or h == "Best app":
                break
            rep["apps"].append(B.canonical_app(h))
    if not rep["apps"]:
        rep["apps"] = B.app_order({app for _, app in rep["runs"]})
    return rep


def report_medians(rep, module=None):
    """{app: {metric: median or score}} from the workbook, the same way build_report.medians reads runs.csv."""
    out = {}
    for (mid, app), cell in rep["runs"].items():
        nums = [v for v in (_num(x) for x in cell["runs"]) if v is not None]
        if nums:
            out.setdefault(app, {})[mid] = round(statistics.median(nums), 2)
    for (mid, app), cell in rep["facts"].items():
        if _num(cell["value"]) is not None:
            out.setdefault(app, {})[mid] = cell["value"]
    for base, items in rep["items"].items():
        per_app = {}
        for answers in items.values():
            for app, cell in answers.items():
                if isinstance(cell["value"], str) and cell["value"].strip():
                    per_app.setdefault(app, []).append(cell["value"].strip())
        for app, answers in per_app.items():
            score = R.item_score(base, answers)
            if score is not None:
                out.setdefault(app, {})[base] = score
    if module is not None and not B.coldstart_on(module):
        for v in out.values():
            for mid in B.COLDSTART_METRICS:
                v.pop(mid, None)
    return out


# ---------------------------------------------------------------- comparing (pure)

def verdict(metric, before, after):
    """better / worse / same / new for RISE, using the metric's direction and the NOISE_PCT band."""
    if before is None:
        return "new", None, None
    delta = round(after - before, 2)
    pct = round(100 * delta / abs(before), 1) if before else None
    if delta == 0 or (pct is not None and abs(pct) < NOISE_PCT):
        return "same", delta, pct
    lower = B.METRICS[metric]["better"] == "Lower"
    return ("better" if (delta < 0) == lower else "worse"), delta, pct


def metric_rank(med, mid, ranked):
    vals = {a: med[a][mid] for a in ranked if mid in med.get(a, {})}
    return R.rank_of(vals, RISE, higher_is_better=B.METRICS[mid]["better"] == "Higher") if RISE in vals else None


def compare(before, after, ranked, weights):
    """before: {app: {mid: v}} from the report; after: {mid: v}, the re-measured RISE medians.

    Returns {"rows": [per metric], "kept": [RISE metrics not re-measured], "overall": (b, a), "rank": (b, a),
             "ranked_of": (b, a), "categories": {cat: (b, a)}}. Competitor values are the report's own."""
    merged = {a: dict(v) for a, v in before.items()}
    merged.setdefault(RISE, {}).update(after)
    rows = []
    for m in B.METRIC_CATALOGUE:
        mid = m["id"]
        if mid not in after:
            continue
        b = before.get(RISE, {}).get(mid)
        v, delta, pct = verdict(mid, b, after[mid])
        rows.append({"id": mid, "category": m["category"], "unit": m["unit"], "better": m["better"],
                     "before": b, "after": after[mid], "delta": delta, "pct": pct, "verdict": v,
                     "rank_before": metric_rank(before, mid, ranked), "rank_after": metric_rank(merged, mid, ranked),
                     "apps": sum(1 for a in ranked if mid in merged.get(a, {}))})
    kept = [mid for mid in (m["id"] for m in B.METRIC_CATALOGUE) if mid in before.get(RISE, {}) and mid not in after]

    def scores(med):
        cs = R.category_scores(R.metric_scores(med, ranked), ranked)
        ov = R.overall(cs, weights, ranked)
        return cs, ov

    cs_b, ov_b = scores(before)
    cs_a, ov_a = scores(merged)
    cats = {cat: (cs_b.get(cat, {}).get(RISE), cs_a.get(cat, {}).get(RISE)) for cat in B.CATEGORIES
            if RISE in cs_b.get(cat, {}) or RISE in cs_a.get(cat, {})}
    return {"rows": rows, "kept": kept, "overall": (ov_b.get(RISE), ov_a.get(RISE)),
            "rank": (R.rank_of(ov_b, RISE), R.rank_of(ov_a, RISE)),
            "ranked_of": (sum(1 for x in ov_b.values() if x is not None), sum(1 for x in ov_a.values() if x is not None)),
            "categories": cats}


# ---------------------------------------------------------------- new RISE runs

def find_verify_folder(module_name, developer, date=None, device=None):
    """results/_verify/<date>/<module>/<developer>_<device>; the latest date when none is given."""
    dates = [date] if date else sorted((p.name for p in VERIFY_DIR.glob("*") if p.is_dir()), reverse=True)
    for d in dates:
        cands = sorted((VERIFY_DIR / d / module_name).glob(f"{B.safe_developer(developer)}_*"))
        if device:
            exact = B.results_dir(module_name, developer, device, VERIFY_DIR / d)
            if exact in cands:
                return d, exact
        if cands:
            return d, max(cands, key=lambda p: R.latest_ts(B.read_runs(p / "runs.csv")))
    return None, None


def rise_rows(folder):
    """RISE rows from the re-check folder and the _app folder beside it."""
    app_folder = folder.parent.parent / B.APP_SCOPE / folder.name
    rows = [r for r in B.read_runs(folder / "runs.csv") if r["app"] == RISE]
    app_rows = [r for r in B.read_runs(app_folder / "runs.csv") if r["app"] == RISE]
    return rows, app_rows


def fmt(v):
    if v is None:
        return "–"
    if isinstance(v, float):
        return f"{v:g}"
    return str(v)


def signed(v, suffix=""):
    return "–" if v is None else f"{'+' if v > 0 else ''}{v:g}{suffix}"


# ---------------------------------------------------------------- writing

def render_md(module, rep, res, meta):
    rows = res["rows"]
    count = {k: sum(1 for r in rows if r["verdict"] == k) for k in ("better", "worse", "same", "new")}
    (ob, oa), (rb, ra), (nb, na) = res["overall"], res["rank"], res["ranked_of"]
    lines = [f"# RISE re-check — {module.get('title', module['name'])}", "",
             "| | |", "|---|---|",
             f"| **Report compared** | `{meta['report']}` |",
             f"| **Phone** | {rep['device']} (re-check on {meta['device']}) |",
             f"| **RISE version** | {fmt(meta['version_before'])} ({fmt(meta['date_before'])}) → {fmt(meta['version_after'])} ({meta['date']}) |",
             f"| **New runs** | `{meta['folder']}` |",
             f"| **Competitors** | values from the report, not measured again: {', '.join(a for a in rep['apps'] if a != RISE)} |",
             "", "## Headline", "",
             f"- **Overall score:** {fmt(ob)} → {fmt(oa)} ({signed(None if ob is None or oa is None else round(oa - ob, 1), ' pts')})",
             f"- **Overall rank:** {fmt(rb)} of {nb} → {fmt(ra)} of {na}",
             f"- **Metrics re-measured:** {len(rows)} — better {count['better']}, worse {count['worse']}, "
             f"no change (within ±{NOISE_PCT:g} %) {count['same']}, new {count['new']}",
             f"- **Not re-measured (report values kept):** {', '.join(res['kept']) or 'none'}"]
    best = sorted((r for r in rows if r["verdict"] == "better" and r["pct"] is not None), key=lambda r: -abs(r["pct"]))[:5]
    worst = sorted((r for r in rows if r["verdict"] == "worse" and r["pct"] is not None), key=lambda r: -abs(r["pct"]))[:5]
    if best:
        lines += ["", "**Biggest improvements:** " + "; ".join(
            f"{r['id']} {B.render(B.METRICS[r['id']]['name'], module)} {fmt(r['before'])} → {fmt(r['after'])} {r['unit']} ({signed(r['pct'], ' %')})" for r in best)]
    if worst:
        lines += ["", "**Regressions:** " + "; ".join(
            f"{r['id']} {B.render(B.METRICS[r['id']]['name'], module)} {fmt(r['before'])} → {fmt(r['after'])} {r['unit']} ({signed(r['pct'], ' %')})" for r in worst)]
    lines += ["", "## Category scores (RISE, 0–100)", "", "| Category | Before | After | Change |", "|---|---|---|---|"]
    for cat, (b, a) in res["categories"].items():
        lines.append(f"| {cat} | {fmt(b)} | {fmt(a)} | {signed(None if b is None or a is None else round(a - b, 1))} |")
    lines += ["", "## Per metric", "",
              "| ID | Metric | Unit | Better | Before | After | Change | Change % | Result | RISE rank before → after |",
              "|---|---|---|---|---|---|---|---|---|---|"]
    for r in rows:
        lines.append(f"| {r['id']} | {B.render(B.METRICS[r['id']]['name'], module)} | {r['unit']} | {r['better']} | {fmt(r['before'])} | "
                     f"{fmt(r['after'])} | {signed(r['delta'])} | {signed(r['pct'], ' %')} | {r['verdict']} | "
                     f"{fmt(r['rank_before'])} → {fmt(r['rank_after'])} of {r['apps']} |")
    lines += ["", "## How to read this", "",
              f"- Before = the report's RISE median; after = the median of the new runs. A change under ±{NOISE_PCT:g} % is "
              "shown as no change, because run-to-run noise on one phone is about that size.",
              "- Scores and ranks are recomputed with the new RISE values against the competitors' values from the "
              "report date. Competitor apps may have changed since then; measure them again before quoting a rank as current.",
              "- Same phone model as the report. Raw values from different phones are never compared.", ""]
    return "\n".join(lines)


def update_workbook(wb, rep, res, rows, app_rows, meta):
    """Write the re-measured RISE values into the workbook's RISE cells and add the verification sheet.
    Returns the list of metrics written. Competitor cells are not touched."""
    written = []
    measured = {r["id"]: r for r in res["rows"]}
    latest = R.latest_per_run(rows + app_rows)
    for mid, r in measured.items():
        if mid in ("O1", "D1"):
            continue
        cell = rep["runs"].get((mid, RISE))
        if cell:
            ws = wb[cell["sheet"]]
            vals = {int(float(x.get("run") or 0)): x["num"] for x in latest if x["metric"] == mid and x["num"] is not None}
            for k in range(B.NRUNS):
                ws.cell(row=cell["row"], column=R.RUN_FIRST + k).value = vals.get(k + 1)
            ws.cell(row=cell["row"], column=R.COL_NOTE).value = (
                f"Re-check {meta['date']}, RISE {fmt(meta['version_after'])}: was median {fmt(r['before'])}")
            written.append(mid)
            continue
        fact = rep["facts"].get((mid, RISE))
        if fact:
            wb[fact["sheet"]][fact["cell"]].value = r["after"]
            written.append(mid)
    for base in ("O1", "D1"):
        if base not in measured:
            continue
        for x in latest:
            if x["metric"].startswith(base + "."):
                cell = rep["items"][base].get(x["metric"].split(".", 1)[1], {}).get(RISE)
                if cell:
                    wb[cell["sheet"]][cell["cell"]].value = (x.get("value") or "").strip()
        written.append(base)
    info = rep["info"].get(RISE)
    if info and meta["version_after"]:
        wb[B.INFO].cell(row=info["row"], column=3).value = meta["version_after"]
        wb[B.INFO].cell(row=info["row"], column=4).value = meta["date"]

    if VERIFY_SHEET in wb.sheetnames:
        del wb[VERIFY_SHEET]
    ws = wb.create_sheet(VERIFY_SHEET, wb.sheetnames.index("Scorecard") + 1 if "Scorecard" in wb.sheetnames else None)
    (ob, oa), (rb, ra), (nb, na) = res["overall"], res["rank"], res["ranked_of"]
    ws["A1"] = f"RISE re-check on {meta['date']}: RISE {fmt(meta['version_before'])} → {fmt(meta['version_after'])}"
    ws["A2"] = (f"Overall {fmt(ob)} → {fmt(oa)}, rank {fmt(rb)} of {nb} → {fmt(ra)} of {na}. "
                f"Competitor values are from the original report. Changes under ±{NOISE_PCT:g} % count as no change.")
    head = ["ID", "Metric", "Unit", "Better", "Before", "After", "Change", "Change %", "Result", "Rank before", "Rank after"]
    for i, h in enumerate(head):
        ws.cell(row=4, column=1 + i, value=h)
    for j, r in enumerate(res["rows"]):
        vals = [r["id"], B.render(B.METRICS[r["id"]]["name"], meta["module"]), r["unit"], r["better"], r["before"], r["after"],
                r["delta"], r["pct"], r["verdict"], r["rank_before"], r["rank_after"]]
        for i, v in enumerate(vals):
            ws.cell(row=5 + j, column=1 + i, value=v)
    ws.column_dimensions["B"].width = 60
    return written


# ---------------------------------------------------------------- commands

def load_report(path):
    from openpyxl import load_workbook
    path = Path(path)
    if path.suffix.lower() != ".xlsx" or not path.exists():
        sys.exit(f"{path}: the re-check needs the module workbook (.xlsx) from build_report.py.")
    wb = load_workbook(path)
    return wb, parse_workbook(wb)


def module_from_report(path, name=None):
    name = name or Path(path).name.split("_", 1)[0]
    return B.load_module(name)


def cmd_baseline(a):
    module = module_from_report(a.report, a.module)
    _, rep = load_report(a.report)
    med = report_medians(rep, module)
    rise = med.get(RISE, {})
    info = rep["info"].get(RISE, {})
    print(f"  Report:  {a.report}\n  Module:  {module['name']}\n  Phone:   {rep['device']}\n"
          f"  RISE:    {fmt(info.get('version'))} tested {fmt(info.get('date'))}\n  Apps:    {', '.join(rep['apps'])}")
    if not rise:
        sys.exit("  The report has no RISE values to compare against.")
    print(f"\n  RISE values in the report ({len(rise)}):")
    for m in B.METRIC_CATALOGUE:
        if m["id"] in rise:
            how = "bench.py" if m["id"] in B.AUTOMATED else "manual"
            print(f"    {m['id']:<4} {fmt(rise[m['id']]):>10} {m['unit']:<6} {how:<8} {B.render(m['name'], module)}")
    groups = [g for g, ids in B.BENCH_GROUPS.items() if any(i in rise for i in ids)]
    mod_groups = [g for g in groups if g not in ("coldstart",)]
    app_groups = [g for g in groups if g in ("coldstart", "ram") and any(i in rise for i in B.BENCH_GROUPS[g] if B.METRICS[i]["phone"])]
    manual = [m for m in rise if m not in B.AUTOMATED]
    print("\n  To measure RISE again (results go to results/_verify/<date>):")
    if app_groups:
        print(f"    _app groups:   {','.join(app_groups)}")
    if mod_groups:
        extra = [f for f, mid in (("--with-r3", "R3"), ("--with-video", "U5")) if mid in rise]
        print(f"    module groups: {','.join(mod_groups)}" + (f" (+ {' '.join(extra)})" if extra else ""))
    if manual:
        print(f"    manual (record.py): {', '.join(sorted(manual))}")


def cmd_compare(a):
    module = module_from_report(a.report, a.module)
    developer = B.get_developer(a.developer)
    wb, rep = load_report(a.report)
    date, folder = find_verify_folder(module["name"], developer, a.date, rep["device"])
    if not folder:
        sys.exit(f"No re-check runs for {developer} under {VERIFY_DIR}/{a.date or '<date>'}/{module['name']}/. "
                 f"Run: python3 bench.py run RISE --module {module['name']} --results-dir results/_verify/<date>")
    device = folder.name.split("_", 1)[1]
    if B.safe_name(device) != B.safe_name(rep["device"] or ""):
        sys.exit(f"The re-check ran on {device}, the report on {rep['device']}. Raw values from different phones are not "
                 "comparable; run the re-check on the report's phone.")
    rows, app_rows = rise_rows(folder)
    app_rows = R.module_app_rows(module, app_rows)
    if not rows and not app_rows:
        sys.exit(f"{folder} has no RISE rows yet.")
    after = R.medians(rows + app_rows, module, {RISE}).get(RISE, {})
    before = report_medians(rep, module)
    ranked = [x for x in rep["apps"] if B.load_apps().get(x, {}).get("ranked", True)]
    res = compare(before, after, ranked, B.load_weights())
    newest = max(rows + app_rows, key=lambda r: r.get("timestamp") or "")
    info = rep["info"].get(RISE, {})
    shown = folder.relative_to(B.HERE) if folder.is_relative_to(B.HERE) else folder
    meta = {"report": a.report, "device": device, "date": date, "folder": str(shown),
            "version_before": info.get("version"), "date_before": info.get("date"),
            "version_after": newest.get("version"), "module": module}

    out_dir = Path(a.out) if a.out else B.REPORTS_DIR / module["name"]
    out_dir.mkdir(parents=True, exist_ok=True)
    md = out_dir / f"verify_{module['name']}_{folder.name}_{date}.md"
    md.write_text(render_md(module, rep, res, meta))
    count = {k: sum(1 for r in res["rows"] if r["verdict"] == k) for k in ("better", "worse", "same", "new")}
    (ob, oa), (rb, ra), (nb, na) = res["overall"], res["rank"], res["ranked_of"]
    print(f"  RISE {fmt(meta['version_before'])} → {fmt(meta['version_after'])} on {device}")
    print(f"  Overall {fmt(ob)} → {fmt(oa)}; rank {fmt(rb)} of {nb} → {fmt(ra)} of {na}")
    print(f"  {len(res['rows'])} metric(s): better {count['better']}, worse {count['worse']}, same {count['same']}, new {count['new']}")
    for r in res["rows"]:
        print(f"    {r['id']:<4} {fmt(r['before']):>9} → {fmt(r['after']):<9} {r['unit']:<6} {signed(r['pct'], ' %'):>9}  {r['verdict']}")
    if res["kept"]:
        print(f"  Not re-measured (report values kept): {', '.join(res['kept'])}")
    print(f"  Comparison: {md}")

    if a.update:
        path = Path(a.report)
        backup = path.with_name(f"{path.stem}.backup-{datetime.datetime.now():%Y%m%d-%H%M%S}.xlsx")
        shutil.copy2(path, backup)
        written = update_workbook(wb, rep, res, rows, app_rows, meta)
        wb.save(path)
        print(f"  Updated {path} ({len(written)} RISE metric(s): {', '.join(written) or 'none'}); backup {backup.name}")
        missing = [r["id"] for r in res["rows"] if r["id"] not in written]
        if missing:
            print(f"  ! Not in the workbook, so not written: {', '.join(missing)}")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("baseline", help="RISE values in a report and what to measure again")
    b.add_argument("--report", required=True); b.add_argument("--module")
    b.set_defaults(fn=cmd_baseline)
    c = sub.add_parser("compare", help="compare the re-check runs with the report")
    c.add_argument("--report", required=True); c.add_argument("--module")
    c.add_argument("--developer"); c.add_argument("--date", help="re-check folder under results/_verify/ (default: latest)")
    c.add_argument("--out", help="folder for the comparison .md (default reports/<module>)")
    c.add_argument("--update", action="store_true", help="write the new RISE values into the same workbook (backup kept)")
    c.set_defaults(fn=cmd_compare)
    a = p.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
